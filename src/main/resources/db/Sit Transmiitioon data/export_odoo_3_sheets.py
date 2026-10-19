import json
import urllib.request
import time
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_URL = 'https://sit-transmission1.odoo.com'
COOKIE = 'frontend_lang=fr_FR; tz=Africa/Casablanca; session_id=5AozDPwL1wo4495It4BJgS6tpadMkaB-cVkRcScq6ETkGLGf9Ll08Qbs-nVg_AI7qzVNH6blC16og1J7-mSNgQ; cids=1; color_scheme=dark'

headers = {
    'Accept': '*/*',
    'Content-Type': 'application/json',
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/155.0.0.0 Safari/537.36 Edg/155.0.0.0',
    'Cookie': COOKIE,
    'Origin': BASE_URL,
    'Referer': f'{BASE_URL}/odoo',
}

STATE_LABELS = {
    'draft': 'Demande de prix',
    'sent': 'Demande envoyée',
    'to approve': 'À approuver',
    'purchase': 'Bon de commande',
    'done': 'Verrouillé',
    'cancel': 'Annulé'
}

RECEIPT_STATUS_LABELS = {
    'no': 'Non reçu',
    'pending': 'En attente',
    'partial': 'Partiellement reçu',
    'full': 'Totalement reçu',
    False: '-',
    None: '-'
}

INVOICE_STATUS_LABELS = {
    'no': 'Rien à facturer',
    'to invoice': 'À facturer',
    'invoiced': 'Facturé',
    False: '-',
    None: '-'
}

INVOICE_STATE_LABELS = {
    'draft': 'Brouillon',
    'posted': 'Comptabilisé',
    'cancel': 'Annulé'
}

PAYMENT_STATE_LABELS = {
    'not_paid': 'Non payé',
    'in_payment': 'En cours',
    'paid': 'Payé',
    'partial': 'Partiel',
    'reversed': 'Extourné',
    'blocked': 'Bloqué',
    False: '-',
    None: '-'
}

MOVE_TYPE_LABELS = {
    'in_invoice': 'Facture Fournisseur',
    'in_refund': 'Avoir Fournisseur',
    'out_invoice': 'Facture Client',
    'out_refund': 'Avoir Client'
}

STOCK_MOVE_STATE_LABELS = {
    'done': 'Fait (Réalisé)',
    'assigned': 'Prêt',
    'confirmed': 'En attente',
    'partially_available': 'Partiellement dispo',
    'draft': 'Brouillon',
    'cancel': 'Annulé'
}

def classify_stock_move(loc_src, loc_dst, ref):
    loc_src_lower = str(loc_src).lower()
    loc_dst_lower = str(loc_dst).lower()
    ref_upper = str(ref or "").upper()

    if "/IN/" in ref_upper or "vendor" in loc_src_lower or "fournisseur" in loc_src_lower:
        return "Entrée (Réception)"
    elif "/OUT/" in ref_upper or "customer" in loc_dst_lower or "client" in loc_dst_lower:
        return "Sortie (Livraison)"
    elif "/INT/" in ref_upper or ("stock" in loc_src_lower and "stock" in loc_dst_lower):
        return "Transfert Interne"
    elif "inventory" in loc_src_lower or "inventory" in loc_dst_lower or "ajustement" in loc_src_lower or "adjustment" in loc_dst_lower:
        return "Ajustement Inventaire"
    elif "scrap" in loc_src_lower or "scrap" in loc_dst_lower or "rebut" in loc_dst_lower:
        return "Rebut / Perte"
    else:
        return "Mouvement Stock"

def fetch_odoo(model, specification, domain=None, context=None, limit=10000, order="id desc"):
    url = f"{BASE_URL}/web/dataset/call_kw/{model}/web_search_read"
    payload = {
        "id": 100,
        "jsonrpc": "2.0",
        "method": "call",
        "params": {
            "model": model,
            "method": "web_search_read",
            "args": [],
            "kwargs": {
                "specification": specification,
                "offset": 0,
                "order": order,
                "limit": limit,
                "domain": domain if domain is not None else [],
                "context": context if context is not None else {"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1]}
            }
        }
    }
    
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode('utf-8'),
        headers=headers,
        method='POST'
    )
    
    with urllib.request.urlopen(req) as resp:
        res_data = resp.read().decode('utf-8')
        result_json = json.loads(res_data)
        if "error" in result_json:
            print(f"[ERREUR sur {model}]: {result_json['error']}")
            return []
        return result_json.get("result", {}).get("records", [])

def format_val(v):
    if isinstance(v, dict):
        return v.get("display_name", str(v))
    elif isinstance(v, list):
        items = []
        for item in v:
            if isinstance(item, dict):
                items.append(item.get("display_name", str(item)))
            else:
                items.append(str(item))
        return ", ".join(items) if items else ""
    return v

# --- Styles OpenPyXL Professionnels ---
FONT_NAME = "Segoe UI"
header_font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(fill_type="solid", start_color="1F4E79", end_color="1F4E79") # Bleu nuit élégant

po_header_font = Font(name=FONT_NAME, size=10.5, bold=True, color="FFFFFF")
po_header_fill = PatternFill(fill_type="solid", start_color="2A4D69", end_color="2A4D69") # Bleu profond commande

sub_header_font = Font(name=FONT_NAME, size=9.5, bold=True, italic=True, color="1F4E79")
sub_header_fill = PatternFill(fill_type="solid", start_color="D9E1F2", end_color="D9E1F2") # Bleu doux / ardoise clair

data_font = Font(name=FONT_NAME, size=9.5, color="1A202C")
data_font_bold = Font(name=FONT_NAME, size=9.5, bold=True, color="1A202C")
data_font_italic = Font(name=FONT_NAME, size=9.5, italic=True, color="6B7280")

zebra_fill_1 = PatternFill(fill_type="solid", start_color="FFFFFF", end_color="FFFFFF")
zebra_fill_2 = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")
sep_fill = PatternFill(fill_type="solid", start_color="EEF2F6", end_color="EEF2F6")

thin_border_side = Side(style='thin', color='E2E8F0')
dark_border_side = Side(style='thin', color='1F4E79')
po_border_side = Side(style='thin', color='2A4D69')

border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_po_header = Border(left=po_border_side, right=po_border_side, top=po_border_side, bottom=po_border_side)
border_sub_header = Border(left=thin_border_side, right=thin_border_side, top=dark_border_side, bottom=dark_border_side)

inv_header_font = Font(name=FONT_NAME, size=10.5, bold=True, color="FFFFFF")
inv_header_fill = PatternFill(fill_type="solid", start_color="164E63", end_color="164E63") # Cyan profond / Émeraude corporate
inv_border_side = Side(style='thin', color='164E63')
border_inv_header = Border(left=inv_border_side, right=inv_border_side, top=inv_border_side, bottom=inv_border_side)

inv_sub_header_font = Font(name=FONT_NAME, size=9.5, bold=True, italic=True, color="164E63")
inv_sub_header_fill = PatternFill(fill_type="solid", start_color="CFFAFE", end_color="CFFAFE") # Teinte douce cyan/émeraude
border_inv_sub_header = Border(left=thin_border_side, right=thin_border_side, top=inv_border_side, bottom=inv_border_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

def auto_fit_columns(ws, max_cols=None):
    ws.views.sheetView[0].showGridLines = True
    columns = list(ws.columns)
    if max_cols:
        columns = columns[:max_cols]
    for col in columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > 60:
                max_len = max(max_len, 40)
            else:
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def build_detailed_orders_sheet(ws, orders):
    ws.title = "Commandes_Detaillees"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    headers_list = [
        "N° Commande / Code Article",
        "Fournisseur / Désignation Article",
        "Réf Fournisseur / Date Prévue",
        "Date Commande",
        "Statut Commande / Ligne",
        "Statut Réception",
        "Statut Facturation",
        "Acheteur / Responsable",
        "Qté Commandée",
        "Qté Reçue",
        "Qté Facturée",
        "Unité",
        "Prix Unitaire HT",
        "Remise %",
        "Taxes",
        "Sous-Total HT",
        "Total TTC",
        "Devise"
    ]
    
    # 1. Ligne d'entête principale du tableau
    ws.append(headers_list)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers_list) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_header
        cell.border = border_po_header

    current_row = 2

    for po in orders:
        order_name = po.get("name", "")
        partner_name = format_val(po.get("partner_id", ""))
        partner_ref = po.get("partner_ref") or "-"
        date_order = po.get("date_order") or "-"
        state_raw = po.get("state", "")
        state_label = STATE_LABELS.get(state_raw, state_raw)
        receipt_raw = po.get("receipt_status", False)
        receipt_label = RECEIPT_STATUS_LABELS.get(receipt_raw, str(receipt_raw) if receipt_raw else "-")
        inv_raw = po.get("invoice_status", False)
        inv_label = INVOICE_STATUS_LABELS.get(inv_raw, str(inv_raw) if inv_raw else "-")
        user_name = format_val(po.get("user_id", ""))
        amount_untaxed = po.get("amount_untaxed", 0.0) or 0.0
        amount_total = po.get("amount_total", 0.0) or 0.0
        currency = format_val(po.get("currency_id", ""))
        order_lines = po.get("order_line", [])

        # --- LIGNE ENTÊTE DE LA COMMANDE FOURNISSEUR ---
        po_row_values = [
            order_name,
            partner_name,
            f"Réf: {partner_ref}" if partner_ref != "-" else "-",
            date_order,
            state_label,
            receipt_label,
            inv_label,
            user_name,
            sum((l.get("product_qty") or 0.0) for l in order_lines),  # Total Qté
            sum((l.get("qty_received") or 0.0) for l in order_lines), # Total Reçue
            sum((l.get("qty_invoiced") or 0.0) for l in order_lines), # Total Facturée
            "-",
            "TOTAL CMD :",
            "-",
            "-",
            amount_untaxed,
            amount_total,
            currency
        ]
        ws.append(po_row_values)
        ws.row_dimensions[current_row].height = 24
        
        for col_idx in range(1, len(po_row_values) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = po_header_font
            cell.fill = po_header_fill
            cell.border = border_po_header
            if col_idx in (1, 4, 5, 6, 7, 12, 14, 15, 18):
                cell.alignment = align_center
            elif col_idx in (9, 10, 11):
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif col_idx in (16, 17):
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif col_idx == 13:
                cell.alignment = align_right
            else:
                cell.alignment = align_left
        
        current_row += 1

        # Si la commande a des lignes d'articles
        if order_lines:
            # --- SOUS-ENTÊTE DES LIGNES D'ARTICLES ---
            sub_headers = [
                "   ↳ Code Article",
                "   ↳ Désignation Article",
                "Date Prévue",
                "",
                "Statut Ligne",
                "",
                "",
                "",
                "Qté Cde",
                "Qté Reçue",
                "Qté Fact.",
                "Unité",
                "Prix Unit. HT",
                "Remise %",
                "Taxes",
                "Sous-Total HT",
                "Total TTC",
                "Devise"
            ]
            ws.append(sub_headers)
            ws.row_dimensions[current_row].height = 20
            for col_idx in range(1, len(sub_headers) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = sub_header_font
                cell.fill = sub_header_fill
                cell.border = border_sub_header
                if col_idx in (3, 5, 12, 15, 18):
                    cell.alignment = align_center
                elif col_idx in (9, 10, 11, 13, 14, 16, 17):
                    cell.alignment = align_right
                else:
                    cell.alignment = align_left
            current_row += 1

            # --- LIGNES ORDRE D'ARTICLES EN BAS DE L'ENTÊTE ---
            for l_idx, line in enumerate(order_lines):
                prod_obj = line.get("product_id")
                prod_name = format_val(prod_obj) if prod_obj else (line.get("name") or "")
                
                # Extraire le code s'il est au format [CODE] Description
                prod_code = ""
                desc_text = line.get("name") or prod_name
                if "[" in prod_name and "]" in prod_name:
                    prod_code = prod_name.split("]")[0].replace("[", "").strip()
                elif "[" in desc_text and "]" in desc_text:
                    prod_code = desc_text.split("]")[0].replace("[", "").strip()
                else:
                    prod_code = f"Ligne #{line.get('sequence', l_idx + 1)}"

                date_planned = line.get("date_planned") or "-"
                line_state = STATE_LABELS.get(line.get("state"), line.get("state") or "-")
                uom_name = format_val(line.get("uom_id", ""))
                p_qty = line.get("product_qty") or 0.0
                q_rec = line.get("qty_received") or 0.0
                q_inv = line.get("qty_invoiced") or 0.0
                p_unit = line.get("price_unit") or 0.0
                discount = (line.get("discount") or 0.0) / 100.0
                p_subtotal = line.get("price_subtotal") or 0.0
                taxes_name = format_val(line.get("tax_ids", ""))
                
                # Calcul approximatif TTC si disponible
                p_total = line.get("price_total") or p_subtotal

                line_row_values = [
                    f"   ↳ {prod_code}",
                    desc_text,
                    date_planned,
                    "",
                    line_state,
                    "",
                    "",
                    "",
                    p_qty,
                    q_rec,
                    q_inv,
                    uom_name,
                    p_unit,
                    discount if discount > 0 else "-",
                    taxes_name if taxes_name else "-",
                    p_subtotal,
                    p_total,
                    currency
                ]
                ws.append(line_row_values)
                ws.row_dimensions[current_row].height = 19
                
                row_fill = zebra_fill_1 if l_idx % 2 == 0 else zebra_fill_2
                for col_idx in range(1, len(line_row_values) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.font = data_font
                    cell.fill = row_fill
                    cell.border = border_data
                    
                    if col_idx in (1, 3, 5, 12, 18):
                        cell.alignment = align_center
                    elif col_idx in (9, 10, 11):
                        cell.alignment = align_right
                        cell.number_format = "#,##0.00"
                    elif col_idx == 13:
                        cell.alignment = align_right
                        cell.number_format = "#,##0.00"
                    elif col_idx == 14:
                        cell.alignment = align_right
                        if discount > 0:
                            cell.number_format = "0.0%"
                    elif col_idx in (16, 17):
                        cell.alignment = align_right
                        cell.number_format = "#,##0.00"
                    elif col_idx == 15:
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left
                current_row += 1
        else:
            # Cas d'une commande sans ligne
            empty_row = [
                "",
                "   (Aucun article / ligne dans cette commande)",
                "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""
            ]
            ws.append(empty_row)
            ws.row_dimensions[current_row].height = 18
            for col_idx in range(1, len(empty_row) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = data_font_italic
                cell.fill = zebra_fill_1
                cell.border = border_data
            current_row += 1

        # Ligne de séparation élégante et discrète entre commandes
        sep_row = [""] * len(headers_list)
        ws.append(sep_row)
        ws.row_dimensions[current_row].height = 6
        for col_idx in range(1, len(sep_row) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = sep_fill
        current_row += 1

    auto_fit_columns(ws)


def build_detailed_invoices_sheet(ws, invoices):
    ws.title = "Factures_Frn_Detaillees"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    headers_list = [
        "N° Facture / Code Article",
        "Fournisseur / Désignation Article",
        "Réf Pièce Frn / Bon Livraison",
        "Date Facture",
        "Date Échéance",
        "Type / Statut Facture",
        "Statut Paiement",
        "N° Cde Achat (BC)",
        "Qté Facturée",
        "Unité",
        "Prix Unit. HT",
        "Remise %",
        "Sous-Total HT",
        "Total Taxes",
        "Total TTC",
        "Reste à Payer",
        "Devise"
    ]
    
    ws.append(headers_list)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers_list) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_header
        cell.border = border_po_header

    current_row = 2

    for inv in invoices:
        inv_name = inv.get("name") or ""
        partner_name = format_val(inv.get("partner_id", ""))
        inv_ref = inv.get("ref") or "-"
        inv_date = inv.get("invoice_date") or "-"
        due_date = inv.get("invoice_date_due") or "-"
        move_type_label = MOVE_TYPE_LABELS.get(inv.get("move_type"), inv.get("move_type") or "")
        state_label = INVOICE_STATE_LABELS.get(inv.get("state"), inv.get("state") or "")
        type_state = f"{move_type_label} ({state_label})"
        payment_status = PAYMENT_STATE_LABELS.get(inv.get("payment_state"), str(inv.get("payment_state")) if inv.get("payment_state") else "-")
        origin = inv.get("invoice_origin") or "-"
        amount_untaxed = inv.get("amount_untaxed", 0.0) or 0.0
        amount_tax = inv.get("amount_tax", 0.0) or 0.0
        amount_total = inv.get("amount_total", 0.0) or 0.0
        amount_residual = inv.get("amount_residual", 0.0) or 0.0
        currency = format_val(inv.get("currency_id", ""))
        inv_lines = inv.get("invoice_line_ids", [])
        
        product_lines = [l for l in inv_lines if not l.get("display_type")]
        total_qty = sum((l.get("quantity") or 0.0) for l in product_lines)

        inv_row_values = [
            inv_name,
            partner_name,
            f"Réf: {inv_ref}" if inv_ref != "-" else "-",
            inv_date,
            due_date,
            type_state,
            payment_status,
            origin,
            total_qty,
            "-",
            "TOTAL FACTURE :",
            "-",
            amount_untaxed,
            amount_tax,
            amount_total,
            amount_residual,
            currency
        ]
        ws.append(inv_row_values)
        ws.row_dimensions[current_row].height = 24
        for col_idx in range(1, len(inv_row_values) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = inv_header_font
            cell.fill = inv_header_fill
            cell.border = border_inv_header
            if col_idx in (1, 4, 5, 6, 7, 8, 10, 12, 17):
                cell.alignment = align_center
            elif col_idx in (9, 13, 14, 15, 16):
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif col_idx == 11:
                cell.alignment = align_right
            else:
                cell.alignment = align_left
        
        current_row += 1

        if inv_lines:
            sub_headers = [
                "   ↳ Code Article",
                "   ↳ Désignation Article",
                "Bon Livraison (BL)",
                "", "", "", "",
                "N° Bon de Cde",
                "Qté Fact.",
                "Unité",
                "Prix Unit. HT",
                "Remise %",
                "Sous-Total HT",
                "",
                "Total TTC",
                "",
                ""
            ]
            ws.append(sub_headers)
            ws.row_dimensions[current_row].height = 20
            for col_idx in range(1, len(sub_headers) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = inv_sub_header_font
                cell.fill = inv_sub_header_fill
                cell.border = border_inv_sub_header
                if col_idx in (1, 2):
                    cell.alignment = align_left
                elif col_idx in (3, 8, 10):
                    cell.alignment = align_center
                elif col_idx in (9, 11, 12, 13, 15):
                    cell.alignment = align_right
                else:
                    cell.alignment = align_center

            current_row += 1

            for line_idx, line in enumerate(inv_lines, start=1):
                prod_obj = line.get("product_id")
                prod_full_name = str(format_val(prod_obj)) if prod_obj else ""
                prod_code = ""
                prod_desc = str(line.get("name") or prod_full_name or "")

                if prod_full_name and "[" in prod_full_name and "]" in prod_full_name:
                    try:
                        prod_code = prod_full_name.split("]")[0].replace("[", "").strip()
                        if not prod_desc or prod_desc == prod_full_name:
                            prod_desc = prod_full_name.split("]", 1)[1].strip()
                    except Exception:
                        prod_code = ""
                elif "[" in prod_desc and "]" in prod_desc:
                    try:
                        prod_code = prod_desc.split("]")[0].replace("[", "").strip()
                    except Exception:
                        prod_code = ""

                qty = line.get("quantity") or 0.0
                uom = format_val(line.get("product_uom_id"))
                price_u = line.get("price_unit") or 0.0
                discount = (line.get("discount") or 0.0) / 100.0
                price_sub = line.get("price_subtotal") or 0.0
                price_tot = line.get("price_total") or 0.0
                bl_ref = line.get("x_studio_bon_livraison") or "-"
                po_ref = format_val(line.get("purchase_order_id")) or "-"

                line_fill = zebra_fill_1 if line_idx % 2 == 1 else zebra_fill_2

                line_row = [
                    f"       {prod_code}" if prod_code else "       -",
                    f"   {prod_desc}",
                    bl_ref,
                    "", "", "", "",
                    po_ref,
                    qty,
                    uom,
                    price_u,
                    discount,
                    price_sub,
                    "",
                    price_tot,
                    "",
                    ""
                ]
                ws.append(line_row)
                ws.row_dimensions[current_row].height = 19
                for col_idx in range(1, len(line_row) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.font = data_font
                    cell.fill = line_fill
                    cell.border = border_data

                    if col_idx in (1, 3, 8, 10):
                        cell.alignment = align_center
                    elif col_idx in (9, 11, 13, 15):
                        cell.alignment = align_right
                        cell.number_format = "#,##0.00"
                    elif col_idx == 12:
                        cell.alignment = align_right
                        if discount > 0:
                            cell.number_format = "0.0%"
                    else:
                        cell.alignment = align_left
                current_row += 1
        else:
            empty_row = [
                "",
                "   (Aucune ligne dans cette facture)",
                "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""
            ]
            ws.append(empty_row)
            ws.row_dimensions[current_row].height = 18
            for col_idx in range(1, len(empty_row) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = data_font_italic
                cell.fill = zebra_fill_1
                cell.border = border_data
            current_row += 1

        sep_row = [""] * len(headers_list)
        ws.append(sep_row)
        ws.row_dimensions[current_row].height = 6
        for col_idx in range(1, len(sep_row) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = sep_fill
        current_row += 1

    auto_fit_columns(ws)


def build_standard_table_sheet(ws, title, headers_list, records, row_formatter):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers_list))}1"
    
    ws.append(headers_list)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers_list) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_header
        cell.border = border_data

    for row_idx, rec in enumerate(records, start=2):
        row_vals, num_formats = row_formatter(rec)
        ws.append(row_vals)
        ws.row_dimensions[row_idx].height = 19
        row_fill = zebra_fill_1 if row_idx % 2 == 0 else zebra_fill_2
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = border_data
            fmt = num_formats.get(col_idx)
            if fmt == "currency" or fmt == "float":
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif fmt == "center":
                cell.alignment = align_center
            elif fmt == "right":
                cell.alignment = align_right
            else:
                cell.alignment = align_left

    auto_fit_columns(ws)


def main():
    print("=" * 60)
    print("EXTRACTEUR ODOO TRANSMISSION - V2 (Détail Commandes & Design)")
    print("=" * 60)

    # 1. Fournisseurs
    print("\n--- 1. Récupération des Fournisseurs (res.partner) ---")
    partner_spec = {
        "name": {}, "is_company": {}, "active": {}, "type": {},
        "email": {}, "phone": {}, "street": {}, "city": {}, "zip": {},
        "country_id": {"fields": {"display_name": {}}},
        "vat": {}, "website": {}, "ref": {},
        "category_id": {"fields": {"display_name": {}}},
        "parent_id": {"fields": {"display_name": {}}},
        "write_date": {}
    }
    fournisseurs_raw = fetch_odoo(
        "res.partner", 
        partner_spec, 
        domain=[["supplier_rank", ">", 0]],
        context={"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1], "res_partner_search_mode": "supplier"}
    )
    print(f"Fournisseurs trouvés : {len(fournisseurs_raw)}")

    # 2. Produits
    print("\n--- 2. Récupération des Produits (product.template) ---")
    product_spec = {
        "name": {}, "default_code": {}, "barcode": {},
        "list_price": {}, "standard_price": {}, 
        "qty_available": {}, "virtual_available": {},
        "incoming_qty": {}, "outgoing_qty": {},
        "nbr_moves_in": {}, "nbr_moves_out": {},
        "product_variant_count": {},
        "categ_id": {"fields": {"display_name": {}}},
        "uom_id": {"fields": {"display_name": {}}},
        "is_storable": {}, "type": {},
        "currency_id": {"fields": {"display_name": {}}},
        "write_date": {}
    }
    produits_raw = fetch_odoo(
        "product.template", 
        product_spec, 
        domain=[],
        context={"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1], "purchase_product_template": 1}
    )
    print(f"Produits trouvés : {len(produits_raw)}")

    # 3. Commandes d'Achat avec leurs lignes complètes imbriquées
    print("\n--- 3. Récupération des Commandes & Lignes de Commandes (purchase.order) ---")
    po_spec = {
        "name": {},
        "partner_ref": {},
        "partner_id": {"fields": {"display_name": {}}},
        "date_order": {},
        "date_planned": {},
        "date_approve": {},
        "origin": {},
        "amount_untaxed": {},
        "amount_total": {},
        "currency_id": {"fields": {"display_name": {}}},
        "state": {},
        "receipt_status": {},
        "invoice_status": {},
        "user_id": {"fields": {"display_name": {}}},
        "company_id": {"fields": {"display_name": {}}},
        "write_date": {},
        "order_line": {
            "fields": {
                "sequence": {},
                "product_id": {"fields": {"display_name": {}}},
                "name": {},
                "product_qty": {},
                "qty_received": {},
                "qty_invoiced": {},
                "uom_id": {"fields": {"display_name": {}}},
                "price_unit": {},
                "discount": {},
                "price_subtotal": {},
                "price_total": {},
                "tax_ids": {"fields": {"display_name": {}}},
                "date_planned": {},
                "state": {}
            },
            "order": "sequence ASC, id ASC"
        }
    }
    commandes_raw = fetch_odoo(
        "purchase.order", 
        po_spec, 
        domain=[],
        context={"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1]}
    )
    total_lignes = sum(len(c.get("order_line", [])) for c in commandes_raw)
    print(f"Commandes d'achat trouvées : {len(commandes_raw)} ({total_lignes} lignes d'articles au total)")

    # 4. Factures Fournisseurs (account.move) avec leurs lignes d'articles et BL
    print("\n--- 4. Récupération des Factures Fournisseurs (account.move) ---")
    inv_spec = {
        "name": {},
        "ref": {},
        "partner_id": {"fields": {"display_name": {}}},
        "invoice_date": {},
        "invoice_date_due": {},
        "move_type": {},
        "state": {},
        "payment_state": {},
        "amount_untaxed": {},
        "amount_tax": {},
        "amount_total": {},
        "amount_residual": {},
        "currency_id": {"fields": {"display_name": {}}},
        "invoice_origin": {},
        "write_date": {},
        "invoice_line_ids": {
            "fields": {
                "sequence": {},
                "display_type": {},
                "product_id": {"fields": {"display_name": {}}},
                "name": {},
                "quantity": {},
                "product_uom_id": {"fields": {"display_name": {}}},
                "price_unit": {},
                "discount": {},
                "price_subtotal": {},
                "price_total": {},
                "tax_ids": {"fields": {"display_name": {}}},
                "purchase_order_id": {"fields": {"display_name": {}}},
                "x_studio_bon_livraison": {}
            },
            "order": "sequence ASC, id ASC"
        }
    }
    factures_raw = fetch_odoo(
        "account.move",
        inv_spec,
        domain=[["move_type", "in", ["in_invoice", "in_refund"]]],
        order="partner_id asc, invoice_date desc, id desc",
        context={"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1], "res_partner_search_mode": "supplier"}
    )
    total_lignes_factures = sum(len(f.get("invoice_line_ids", [])) for f in factures_raw)
    print(f"Factures fournisseurs trouvées : {len(factures_raw)} ({total_lignes_factures} lignes d'articles au total)")

    # 5. Mouvements de Stock (Entrées, Sorties, Réceptions, Livraisons)
    print("\n--- 5. Récupération des Mouvements de Stock (stock.move) ---")
    stock_move_spec = {
        "id": {},
        "date": {},
        "reference": {},
        "product_id": {"fields": {"display_name": {}}},
        "quantity": {},
        "product_uom_qty": {},
        "location_id": {"fields": {"display_name": {}}},
        "location_dest_id": {"fields": {"display_name": {}}},
        "partner_id": {"fields": {"display_name": {}}},
        "origin": {},
        "state": {}
    }
    stock_moves_raw = fetch_odoo(
        "stock.move",
        stock_move_spec,
        domain=[],
        order="date desc, id desc",
        context={"lang": "fr_FR", "tz": "Africa/Casablanca", "uid": 2, "allowed_company_ids": [1]}
    )
    print(f"Mouvements de stock trouvés : {len(stock_moves_raw)}")

    # 6. Génération du classeur principal avec design (7 onglets)
    output_excel_main = "odoo_export_transmission1.xlsx"
    print(f"\n--- 6. Création du classeur Excel principal '{output_excel_main}' avec design ---")
    wb_main = openpyxl.Workbook()
    
    # Onglet 1: Fournisseurs
    ws_fournisseurs = wb_main.active
    fourn_headers = ["Nom / Raison Sociale", "Société ?", "Actif", "Type", "Email", "Téléphone", "Adresse", "Ville", "Code Postal", "Pays", "N° TVA", "Site Web", "Réf Interne", "Étiquettes", "Société Mère", "Dernière MAJ"]
    def format_fournisseur(r):
        vals = [
            r.get("name") or "",
            "Oui" if r.get("is_company") else "Non",
            "Oui" if r.get("active") else "Non",
            r.get("type") or "",
            r.get("email") or "",
            r.get("phone") or "",
            r.get("street") or "",
            r.get("city") or "",
            r.get("zip") or "",
            format_val(r.get("country_id")),
            r.get("vat") or "",
            r.get("website") or "",
            r.get("ref") or "",
            format_val(r.get("category_id")),
            format_val(r.get("parent_id")),
            r.get("write_date") or ""
        ]
        formats = {2: "center", 3: "center", 4: "center", 8: "center", 9: "center", 10: "center", 13: "center", 16: "center"}
        return vals, formats
    build_standard_table_sheet(ws_fournisseurs, "Fournisseurs", fourn_headers, fournisseurs_raw, format_fournisseur)

    # Onglet 2: Produits (avec Variantes, Flux Entrants & Sortants)
    ws_produits = wb_main.create_sheet(title="Produits")
    prod_headers = [
        "Nom de l'Article", "Référence Interne", "Code-barres", 
        "Prix de Vente HT", "Coût Standard (PMP)", 
        "Qté en Stock", "Qté Prévue", "Catégorie", "Unité de Mesure", 
        "Stockable ?", "Type", "Devise", "Dernière MAJ",
        "Qté Entrante (Cde Frn)", "Qté Sortante (Réservée Client)", 
        "Nb Entrées (Réceptions)", "Nb Sorties (Livraisons)", "Nb Variantes"
    ]
    def format_produit(r):
        vals = [
            r.get("name") or "",
            r.get("default_code") or "",
            r.get("barcode") or "",
            r.get("list_price") or 0.0,
            r.get("standard_price") or 0.0,
            r.get("qty_available") or 0.0,
            r.get("virtual_available") or 0.0,
            format_val(r.get("categ_id")),
            format_val(r.get("uom_id")),
            "Oui" if r.get("is_storable") else "Non",
            r.get("type") or "",
            format_val(r.get("currency_id")),
            r.get("write_date") or "",
            r.get("incoming_qty") or 0.0,
            r.get("outgoing_qty") or 0.0,
            r.get("nbr_moves_in") or 0,
            r.get("nbr_moves_out") or 0,
            r.get("product_variant_count") or 1
        ]
        formats = {
            2: "center", 3: "center", 4: "currency", 5: "currency", 
            6: "float", 7: "float", 10: "center", 11: "center", 12: "center", 13: "center",
            14: "float", 15: "float", 16: "center", 17: "center", 18: "center"
        }
        return vals, formats
    build_standard_table_sheet(ws_produits, "Produits", prod_headers, produits_raw, format_produit)

    # Onglet 3: Mouvements de Stock (Historique complet des Entrées, Sorties et Transferts)
    ws_mouvements = wb_main.create_sheet(title="Mouvements_Stock")
    stock_headers = [
        "Date Mouvement", "Type d'Opération", "N° Bon / Pièce", 
        "Code Article", "Désignation Article", 
        "Qté Réalisée", "Qté Demandée", 
        "Emplacement Source", "Emplacement Destination", 
        "Partenaire (Frn / Client)", "Document Origine (BC/SO)", "Statut"
    ]
    def format_stock_move(r):
        loc_src = format_val(r.get("location_id")) or ""
        loc_dst = format_val(r.get("location_dest_id")) or ""
        ref_bon = r.get("reference") or "-"
        op_type = classify_stock_move(loc_src, loc_dst, ref_bon)
        
        prod_obj = r.get("product_id")
        prod_full = format_val(prod_obj) or ""
        p_code = ""
        p_desc = prod_full
        if "[" in prod_full and "]" in prod_full:
            try:
                p_code = prod_full.split("]")[0].replace("[", "").strip()
                p_desc = prod_full.split("]", 1)[1].strip()
            except Exception:
                p_code = ""

        vals = [
            r.get("date") or "",
            op_type,
            ref_bon,
            p_code,
            p_desc,
            r.get("quantity") or 0.0,
            r.get("product_uom_qty") or 0.0,
            loc_src,
            loc_dst,
            format_val(r.get("partner_id")) or "-",
            r.get("origin") or "-",
            STOCK_MOVE_STATE_LABELS.get(r.get("state"), r.get("state") or "")
        ]
        formats = {1: "center", 2: "center", 3: "center", 4: "center", 6: "float", 7: "float", 11: "center", 12: "center"}
        return vals, formats
    build_standard_table_sheet(ws_mouvements, "Mouvements_Stock", stock_headers, stock_moves_raw, format_stock_move)

    # Onglet 4: Commandes Détaillées (Entête Commande + Lignes sous l'entête)
    ws_cmd_detail = wb_main.create_sheet(title="Commandes_Detaillees")
    build_detailed_orders_sheet(ws_cmd_detail, commandes_raw)

    # Onglet 5: Commandes Synthèse (Table plate)
    ws_cmd_synthese = wb_main.create_sheet(title="Commandes_Synthese")
    cmd_syn_headers = ["N° Commande", "Fournisseur", "Réf Fournisseur", "Date Commande", "Date Prévue", "Total HT", "Total TTC", "Devise", "Statut", "Statut Réception", "Statut Facturation", "Acheteur", "Document Source", "Dernière MAJ"]
    def format_cmd_syn(r):
        vals = [
            r.get("name") or "",
            format_val(r.get("partner_id")),
            r.get("partner_ref") or "-",
            r.get("date_order") or "",
            r.get("date_planned") or "",
            r.get("amount_untaxed") or 0.0,
            r.get("amount_total") or 0.0,
            format_val(r.get("currency_id")),
            STATE_LABELS.get(r.get("state"), r.get("state") or ""),
            RECEIPT_STATUS_LABELS.get(r.get("receipt_status"), "-"),
            INVOICE_STATUS_LABELS.get(r.get("invoice_status"), "-"),
            format_val(r.get("user_id")),
            r.get("origin") or "-",
            r.get("write_date") or ""
        ]
        formats = {1: "center", 3: "center", 4: "center", 5: "center", 6: "currency", 7: "currency", 8: "center", 9: "center", 10: "center", 11: "center", 14: "center"}
        return vals, formats
    build_standard_table_sheet(ws_cmd_synthese, "Commandes_Synthese", cmd_syn_headers, commandes_raw, format_cmd_syn)

    # Onglet 6: Factures Fournisseurs Détaillées
    ws_inv_detail = wb_main.create_sheet(title="Factures_Frn_Detaillees")
    build_detailed_invoices_sheet(ws_inv_detail, factures_raw)

    # Onglet 7: Factures Fournisseurs Synthèse
    ws_inv_synthese = wb_main.create_sheet(title="Factures_Frn_Synthese")
    inv_syn_headers = ["N° Facture", "Type Document", "Fournisseur", "Réf Pièce Fournisseur", "Date Facture", "Date Échéance", "Statut Comptable", "Statut Paiement", "Total HT", "Total Taxes", "Total TTC", "Reste à Payer", "Devise", "Origine / N° Commande", "Dernière MAJ"]
    def format_inv_syn(r):
        vals = [
            r.get("name") or "",
            MOVE_TYPE_LABELS.get(r.get("move_type"), r.get("move_type") or ""),
            format_val(r.get("partner_id")),
            r.get("ref") or "-",
            r.get("invoice_date") or "",
            r.get("invoice_date_due") or "",
            INVOICE_STATE_LABELS.get(r.get("state"), r.get("state") or ""),
            PAYMENT_STATE_LABELS.get(r.get("payment_state"), str(r.get("payment_state")) if r.get("payment_state") else "-"),
            r.get("amount_untaxed") or 0.0,
            r.get("amount_tax") or 0.0,
            r.get("amount_total") or 0.0,
            r.get("amount_residual") or 0.0,
            format_val(r.get("currency_id")),
            r.get("invoice_origin") or "-",
            r.get("write_date") or ""
        ]
        formats = {1: "center", 2: "center", 4: "center", 5: "center", 6: "center", 7: "center", 8: "center", 9: "currency", 10: "currency", 11: "currency", 12: "currency", 13: "center", 14: "center", 15: "center"}
        return vals, formats
    build_standard_table_sheet(ws_inv_synthese, "Factures_Frn_Synthese", inv_syn_headers, factures_raw, format_inv_syn)

    wb_main.save(output_excel_main)
    print(f"-> Sauvegardé avec succès: '{output_excel_main}' (fichier UNIQUE consolidé avec 7 onglets)")

    print("\n" + "=" * 60)
    print("EXTRACTION TERMINÉE : UN SEUL FICHIER SPREADSHEET GÉNÉRÉ !")
    print("=" * 60)

if __name__ == "__main__":
    main()

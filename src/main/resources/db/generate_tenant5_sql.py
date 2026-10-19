import openpyxl
import os
import sys
import re
from datetime import datetime

# Configure output encoding for Windows
sys.stdout.reconfigure(encoding='utf-8', errors='backslashreplace')

def sql_escape(val, max_len=None):
    if val is None:
        return "NULL"
    s = str(val).strip()
    if s == "" or s.lower() == "none" or s.lower() == "false":
        return "NULL"
    if max_len and len(s) > max_len:
        s = s[:max_len]
    # Escape single quotes
    s = s.replace("'", "''")
    return f"'{s}'"

def num_val(val, default=0.0):
    if val is None:
        return default
    try:
        s = str(val).replace(',', '.').strip()
        if not s or s.lower() == 'none' or s.lower() == 'false' or s == '-':
            return default
        return float(s)
    except:
        return default

def int_val(val, default=0):
    if val is None:
        return default
    try:
        return int(round(num_val(val, default)))
    except:
        return default

def date_val(val):
    if val is None:
        return "CURRENT_TIMESTAMP"
    s = str(val).strip()
    if not s or s.lower() == 'none' or s == '-':
        return "CURRENT_TIMESTAMP"
    m = re.match(r'^(\d{4}-\d{2}-\d{2})[T ]?(\d{2}:\d{2}:\d{2})?', s)
    if m:
        d = m.group(1)
        t = m.group(2) if m.group(2) else "00:00:00"
        return f"'{d} {t}'"
    return "CURRENT_TIMESTAMP"

def map_statut_commande(odoo_st):
    st = str(odoo_st or '').strip().lower()
    if 'bon de commande' in st:
        return 'VALIDEE'
    elif 'demande envoyée' in st or 'demande envoyee' in st:
        return 'PASSEE'
    elif 'annulé' in st or 'annule' in st:
        return 'ANNULEE'
    elif 'demande de prix' in st:
        return 'BROUILLON'
    return 'VALIDEE'

def map_statut_livraison(odoo_rec):
    rec = str(odoo_rec or '').strip().lower()
    if 'totalement' in rec:
        return 'LIVREE'
    elif 'partiellement' in rec:
        return 'PARTIELLE'
    elif 'annul' in rec:
        return 'ANNULEE'
    return 'EN_ATTENTE'

def map_unite(u_str):
    u = str(u_str or '').lower()
    if 'm²' in u or 'm2' in u:
        return 'M2'
    elif 'm' == u or 'mètre' in u:
        return 'METRE'
    elif 'kg' in u or 'kilo' in u:
        return 'KG'
    elif 'l' == u or 'litre' in u:
        return 'LITRE'
    elif 'carton' in u:
        return 'CARTON'
    elif 'palette' in u:
        return 'PALETTE'
    elif 'paquet' in u:
        return 'PAQUET'
    elif 'rouleau' in u:
        return 'ROULEAU'
    return 'PIECE'

def generate():
    excel_path = os.path.join(
        os.path.dirname(__file__),
        "Sit Transmiitioon data",
        "odoo_export_transmission1.xlsx"
    )
    output_sql_path = os.path.join(
        os.path.dirname(__file__),
        "import_sit_transmission_tenant5.sql"
    )

    print(f"Chargement de {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    TENANT_ID = 5
    TENANT_NAME = "SIT TRANSMISSION SARL"

    sql_lines = []
    sql_lines.append("-- =============================================================================")
    sql_lines.append(f"-- IMPORT ODOO -> TENANT #{TENANT_ID} : {TENANT_NAME}")
    sql_lines.append(f"-- Généré le : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    sql_lines.append("-- Source : odoo_export_transmission1.xlsx")
    sql_lines.append("-- =============================================================================\n")
    sql_lines.append("BEGIN;\n")

    # 0. TRUNCATE DES TABLES ACHATS, STOCKS ET CATALOGUE
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 0. TRUNCATE DES TABLES ACHATS, STOCKS & CATALOGUE (RESTART IDENTITY CASCADE)")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("""
TRUNCATE TABLE 
    lignes_commande,
    commandes,
    lignes_livraison,
    livraisons,
    lignes_facture_achat,
    factures_achat,
    reglements_fournisseur,
    lignes_inventaire,
    inventaires,
    lignes_transfert_stock,
    transferts_stock,
    mouvements_stock,
    stock_qualites,
    stocks,
    variantes_produit,
    produit_images,
    produits,
    fournisseurs,
    categories
RESTART IDENTITY CASCADE;
""")

    # 1. Tenant / Point de Vente
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append(f"-- 1. VALIDATION DU POINT DE VENTE / TENANT #{TENANT_ID}")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append(f"""
INSERT INTO point_de_vente (id, tenant_id, nom_point_de_vente, nom, actif, date_creation)
VALUES ({TENANT_ID}, {TENANT_ID}, '{TENANT_NAME}', '{TENANT_NAME}', TRUE, CURRENT_TIMESTAMP)
ON CONFLICT (id) DO UPDATE 
SET nom_point_de_vente = EXCLUDED.nom_point_de_vente,
    nom = EXCLUDED.nom,
    tenant_id = EXCLUDED.tenant_id;

SELECT setval('point_de_vente_id_seq', GREATEST((SELECT COALESCE(MAX(id), 1) FROM point_de_vente), {TENANT_ID}));
""")

    # 2. Catégories
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 2. CATÉGORIES D'ARTICLES")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    ws_prod = wb['Produits']
    categories = set()
    for r in list(ws_prod.iter_rows(values_only=True))[1:]:
        c = str(r[7] or '').strip()
        if c and c.lower() not in ['false', 'none', '']:
            categories.add(c)

    print(f"Catégories trouvées : {len(categories)}")
    for cat in sorted(categories):
        escaped_cat = cat[:255].replace("'", "''")
        code = cat[:10].upper().replace(" ", "_").replace("'", "_")
        sql_lines.append(
            f"INSERT INTO categories (nom, code, actif, point_de_vente_id, date_creation) "
            f"VALUES ('{escaped_cat}', '{code}', TRUE, {TENANT_ID}, CURRENT_TIMESTAMP) "
            f"ON CONFLICT DO NOTHING;"
        )
    sql_lines.append("")

    # 3. Fournisseurs (Classification NATIONAL / INTERNATIONAL & DEVISES)
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 3. FOURNISSEURS (Classification NATIONAL / INTERNATIONAL & DEVISES)")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    ws_frn = wb['Fournisseurs']
    ws_cmd = wb['Commandes_Synthese']

    # Récupérer les devises utilisées par chaque fournisseur dans les commandes
    cmd_devises = {}
    for r in list(ws_cmd.iter_rows(values_only=True))[1:]:
        nom_c = str(r[1] or '').strip()
        dev_c = str(r[7] or 'MAD').strip()
        if nom_c not in cmd_devises:
            cmd_devises[nom_c] = set()
        cmd_devises[nom_c].add(dev_c)

    fournisseurs = {}
    for r in list(ws_frn.iter_rows(values_only=True))[1:]:
        nom = str(r[0] or '').strip()
        if nom and nom.lower() not in ['false', 'none', '']:
            p = str(r[9] or '').strip()
            pays_val = p if p.lower() not in ['false', 'none'] else ''
            fournisseurs[nom] = {
                'adresse': str(r[6] or '').strip() if r[6] and str(r[6]).lower() != 'none' else '',
                'ville': str(r[7] or '').strip() if r[7] and str(r[7]).lower() != 'none' else '',
                'cp': str(r[8] or '').strip() if r[8] and str(r[8]).lower() != 'none' else '',
                'pays': pays_val,
                'ice': str(r[10] or '').strip() if r[10] and str(r[10]).lower() != 'none' else '',
                'email': str(r[4] or '').strip() if r[4] and str(r[4]).lower() != 'none' else '',
                'telephone': str(r[5] or '').strip() if r[5] and str(r[5]).lower() != 'none' else '',
            }

    # Add missing suppliers from Commandes_Synthese
    for r in list(ws_cmd.iter_rows(values_only=True))[1:]:
        nom = str(r[1] or '').strip()
        if nom and nom not in fournisseurs:
            fournisseurs[nom] = {
                'adresse': '', 'ville': '', 'cp': '', 'pays': '',
                'ice': '', 'email': '', 'telephone': ''
            }

    nat_count = 0
    inter_count = 0

    print(f"Fournisseurs totaux à insérer : {len(fournisseurs)}")
    for nom, data in sorted(fournisseurs.items()):
        pays = data.get('pays', '')
        devs = cmd_devises.get(nom, {'MAD'})

        is_international = False
        default_devise = 'MAD'

        if pays and pays.lower() != 'maroc':
            is_international = True
        if 'USD' in devs:
            is_international = True
            default_devise = 'USD'
        elif 'EUR' in devs:
            is_international = True
            default_devise = 'EUR'

        if is_international:
            type_frn = 'INTERNATIONAL'
            inter_count += 1
            if not pays:
                if 'USD' in devs:
                    pays = 'Chine' if any(w in nom.lower() for w in ['ningbo', 'shanghai', 'tech', 'filtration', 'co', 'ltd']) else 'International'
                else:
                    pays = 'Europe' if any(w in nom.lower() for w in ['srl', 'sl', 'tecnico', 'export']) else 'International'
        else:
            type_frn = 'NATIONAL'
            nat_count += 1
            if not pays:
                pays = 'Maroc'
            if not data.get('ville'):
                data['ville'] = 'Casablanca'

        e_nom = nom[:255].replace("'", "''")
        e_adr = sql_escape(data['adresse']) # TEXT in postgres
        e_vil = sql_escape(data['ville'], 100)
        e_cp  = sql_escape(data['cp'], 20)
        e_pay = sql_escape(pays, 100)
        e_ice = sql_escape(data['ice'], 50)
        e_em  = sql_escape(data['email'], 150)
        e_tel = sql_escape(data['telephone'], 50)

        sql_lines.append(
            f"INSERT INTO fournisseurs (raison_social, adresse, ville, code_postal, pays, ice, email, telephone, "
            f"point_de_vente_id, actif, type_fournisseur, devise, date_creation) "
            f"VALUES ('{e_nom}', {e_adr}, {e_vil}, {e_cp}, {e_pay}, {e_ice}, {e_em}, {e_tel}, "
            f"{TENANT_ID}, TRUE, '{type_frn}', '{default_devise}', CURRENT_TIMESTAMP);"
        )

    print(f"Classification des fournisseurs : {nat_count} Nationaux, {inter_count} Internationaux.")
    sql_lines.append("")

    # 4. Produits & Stocks
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 4. PRODUITS & STOCKS INITIALS")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    prod_ref_counts = {}
    prod_map_for_details = {}

    prod_rows = list(ws_prod.iter_rows(values_only=True))[1:]
    print(f"Produits à traiter : {len(prod_rows)}")

    for idx, r in enumerate(prod_rows, start=1):
        raw_nom = str(r[0] or f"Article #{idx}").strip()
        raw_ref = str(r[1] or "").strip()
        cb = str(r[2] or "").strip()
        if not cb or cb.lower() == 'none':
            cb = None
        pv = num_val(r[3], 0.0)
        pa = num_val(r[4], 0.0)
        qte_stock = max(0.0, num_val(r[5], 0.0))
        cat_nom = str(r[7] or "").strip()
        unite = map_unite(r[8])

        # Generate or sanitize reference
        if not raw_ref or raw_ref.lower() == 'none':
            unique_ref = f"ART-ST-{idx:05d}"
        else:
            c = prod_ref_counts.get(raw_ref, 0)
            if c > 0:
                unique_ref = f"{raw_ref}_{c+1}"
            else:
                unique_ref = raw_ref
            prod_ref_counts[raw_ref] = c + 1

        # Store lookup mapping for order lines
        clean_name = re.sub(r'^\[.*?\]\s*', '', raw_nom).strip().lower()
        if raw_ref and raw_ref.lower() != 'none':
            prod_map_for_details[raw_ref] = unique_ref
            prod_map_for_details[(raw_ref, clean_name)] = unique_ref
        prod_map_for_details[clean_name] = unique_ref

        # Truncate designation to 255 chars, keep full raw_nom in description (TEXT)
        e_ref = unique_ref[:255].replace("'", "''")
        e_designation = raw_nom[:255].replace("'", "''")
        e_description = raw_nom.replace("'", "''")
        e_cb = sql_escape(cb, 255)
        e_cat_nom = sql_escape(cat_nom[:255] if cat_nom.lower() not in ['false', 'none', ''] else None, 255)

        if cat_nom and cat_nom.lower() not in ['false', 'none', '']:
            cat_subquery = f"(SELECT id FROM categories WHERE nom = '{cat_nom[:255].replace('\'', '\'\'')}' AND point_de_vente_id = {TENANT_ID} LIMIT 1)"
        else:
            cat_subquery = "NULL"

        # Insert Produit
        sql_lines.append(
            f"INSERT INTO produits (reference, designation, description, code_barre, categorie_id, categorie_article, "
            f"prix_achat_ht, prix_achat_ttc, prix_vente_ht, prix_vente_ttc, prix_achat, prix_vente, "
            f"unite_mesure_stock, stock_minimum, point_de_vente_id, actif, date_creation) "
            f"VALUES ('{e_ref}', '{e_designation}', '{e_description}', {e_cb}, {cat_subquery}, {e_cat_nom}, "
            f"{pa:.2f}, {pa:.2f}, {pv:.2f}, {pv:.2f}, {pa:.2f}, {pv:.2f}, "
            f"'{unite}', 0.00, {TENANT_ID}, TRUE, CURRENT_TIMESTAMP);"
        )

        # Insert Stock
        sql_lines.append(
            f"INSERT INTO stocks (produit_id, quantite_disponible, quantite_reservee, seuil_alerte, derniere_maj) "
            f"VALUES ((SELECT id FROM produits WHERE reference = '{e_ref}' AND point_de_vente_id = {TENANT_ID} LIMIT 1), "
            f"{qte_stock:.2f}, 0.00, 0.00, CURRENT_TIMESTAMP);"
        )

    sql_lines.append("")

    # 5. Commandes
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 5. COMMANDES FOURNISSEURS (218 Bons Odoo)")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    cmd_rows = list(ws_cmd.iter_rows(values_only=True))[1:]
    print(f"Commandes à traiter : {len(cmd_rows)}")

    for r in cmd_rows:
        num_cmd = str(r[0] or "").strip()
        frn_nom = str(r[1] or "").strip()
        ref_frn = str(r[2] or "").strip()
        dt_cmd = date_val(r[3])
        dt_prev = date_val(r[4])
        total_ht = num_val(r[5], 0.0)
        total_ttc = num_val(r[6], 0.0)
        devise = str(r[7] or "MAD").strip()
        st_cmd = map_statut_commande(r[8])
        st_liv = map_statut_livraison(r[9])

        e_num = num_cmd[:100].replace("'", "''")
        e_frn = frn_nom[:255].replace("'", "''")
        e_devise = devise[:10].replace("'", "''")

        montant_tva = max(0.0, total_ttc - total_ht)
        appliquer_tva = "TRUE" if montant_tva > 0.01 else "FALSE"
        taux_tva = "20.00" if montant_tva > 0.01 else "0.00"

        sql_lines.append(
            f"INSERT INTO commandes (numero_commande, fournisseur_id, statut, statut_livraison, "
            f"date_commande, date_livraison_prevue, montant_total, montant_ht, montant_tva, montant_ttc, "
            f"appliquer_tva, taux_tva, devise, point_de_vente_id) "
            f"VALUES ('{e_num}', "
            f"(SELECT id FROM fournisseurs WHERE raison_social = '{e_frn}' AND point_de_vente_id = {TENANT_ID} LIMIT 1), "
            f"'{st_cmd}', '{st_liv}', {dt_cmd}, {dt_prev}, {total_ttc:.2f}, {total_ht:.2f}, {montant_tva:.2f}, {total_ttc:.2f}, "
            f"{appliquer_tva}, {taux_tva}, '{e_devise}', {TENANT_ID});"
        )

    sql_lines.append("")

    # 6. Lignes de Commande
    sql_lines.append("-- -----------------------------------------------------------------------------")
    sql_lines.append("-- 6. LIGNES DE COMMANDE DÉTAILLÉES (967 Lignes)")
    sql_lines.append("-- -----------------------------------------------------------------------------")
    ws_det = wb['Commandes_Detaillees']
    current_cmd = None
    lignes_count = 0

    for r in list(ws_det.iter_rows(values_only=True))[1:]:
        c0 = str(r[0] or "").strip()
        c1 = str(r[1] or "").strip()

        # If it's a command header
        if c0 and not c0.startswith('↳') and not 'Code Article' in c0:
            current_cmd = c0
            continue

        # If it's a detail line
        if '↳' in c0:
            if 'Code Article' in c0:
                continue
            if not current_cmd:
                continue

            raw_code = c0.replace('↳', '').strip()
            clean_name = re.sub(r'^\[.*?\]\s*', '', c1).strip().lower()

            # Find matching product unique_ref
            unique_ref = prod_map_for_details.get((raw_code, clean_name))
            if not unique_ref:
                unique_ref = prod_map_for_details.get(raw_code)
            if not unique_ref:
                unique_ref = prod_map_for_details.get(clean_name)

            if not unique_ref:
                unique_ref = raw_code

            qte_cde = int_val(r[8], 1)
            qte_rec = int_val(r[9], 0)
            pu_ht = num_val(r[12], 0.0)
            st_ht = num_val(r[15], pu_ht * qte_cde)
            st_ttc = num_val(r[16], st_ht)
            montant_tva = max(0.0, st_ttc - st_ht)
            taux_tva = 20.00 if montant_tva > 0.01 else 0.00

            e_cmd = current_cmd[:100].replace("'", "''")
            e_ref = unique_ref[:255].replace("'", "''")

            sql_lines.append(
                f"INSERT INTO lignes_commande (commande_id, produit_id, quantite_commandee, quantite_livree, "
                f"prix_unitaire, prix_unitaire_ht, taux_tva, montant_ht, montant_tva, montant_ttc, montant_ligne, qualite_produit) "
                f"VALUES ("
                f"(SELECT id FROM commandes WHERE numero_commande = '{e_cmd}' AND point_de_vente_id = {TENANT_ID} LIMIT 1), "
                f"(SELECT id FROM produits WHERE reference = '{e_ref}' AND point_de_vente_id = {TENANT_ID} LIMIT 1), "
                f"{qte_cde}, {qte_rec}, {pu_ht:.4f}, {pu_ht:.4f}, {taux_tva:.2f}, {st_ht:.2f}, {montant_tva:.2f}, {st_ttc:.2f}, {st_ht:.2f}, 'PREMIERE_QUALITE');"
            )
            lignes_count += 1

    print(f"Lignes de commande générées : {lignes_count}")
    sql_lines.append("\nCOMMIT;\n")

    print(f"Écriture du fichier SQL : {output_sql_path}...")
    with open(output_sql_path, "w", encoding="utf-8") as f_out:
        f_out.write("\n".join(sql_lines))

    file_size_mb = os.path.getsize(output_sql_path) / (1024 * 1024)
    print(f"Succès ! Fichier SQL généré ({file_size_mb:.2f} MB, {len(sql_lines)} lignes).")

if __name__ == "__main__":
    generate()

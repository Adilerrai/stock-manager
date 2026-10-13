import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_onboarding_workbook(output_filename="MODELE_IMPORT_ONBOARDING_ERP.xlsx"):
    wb = openpyxl.Workbook()
    # Supprimer la feuille par defaut
    default_sheet = wb.active
    
    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Bleu Nuit / Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    sheets_config = [
        {
            "name": "CLIENTS",
            "headers": [
                "code_client", "nom_raison_sociale", "prenom", "forme_juridique",
                "telephone", "email", "adresse", "ville", "pays",
                "ice", "rc", "identifiant_fiscal", "patente",
                "categorie", "tarif", "delai_paiement_jours", "remise_defaut_pct"
            ],
            "examples": [
                [
                    "CLI-0001", "STE ATLAS TRAVAUX SARL", "", "SARL",
                    "0522112233", "contact@atlastravaux.ma", "123 Bd Zerktouni", "Casablanca", "Maroc",
                    "001523489000045", "145892", "40125896", "34567890",
                    "PROFESSIONNEL", "GROS", 60, 5.0
                ],
                [
                    "CLI-0002", "EL ALAMI", "Karim", "PARTICULIER",
                    "0661223344", "karim.alami@gmail.com", "45 Rue Hassan II", "Rabat", "Maroc",
                    "", "", "", "",
                    "PARTICULIER", "DETAIL", 0, 0.0
                ]
            ]
        },
        {
            "name": "FOURNISSEURS",
            "headers": [
                "code_fournisseur", "raison_sociale", "contact_nom", "telephone",
                "email", "adresse", "ville", "pays",
                "ice", "rc", "identifiant_fiscal", "patente",
                "type_fournisseur", "devise", "delai_paiement_jours",
                "banque_nom", "rib_bancaire"
            ],
            "examples": [
                [
                    "FRN-0001", "MAGHREB DISTRIB CHIMIE", "M. Berrada", "0522889900",
                    "commandes@mdc.ma", "Zone Industrielle Ain Sebaa", "Casablanca", "Maroc",
                    "002485698000012", "98542", "30254100", "22558899",
                    "LOCAL", "MAD", 60,
                    "Attijariwafa Bank", "007780001234567890123456"
                ],
                [
                    "FRN-0002", "EUROPE PHARMA LABS", "Sarah Dupont", "+33140506070",
                    "export@europharmalabs.com", "15 Av Montaigne", "Paris", "France",
                    "", "", "", "",
                    "INTERNATIONAL", "EUR", 30,
                    "BNP Paribas", "FR7630004000123456789012345"
                ]
            ]
        },
        {
            "name": "CATEGORIES",
            "headers": [
                "code_categorie", "nom_categorie", "code_categorie_parente",
                "description", "couleur_hex"
            ],
            "examples": [
                ["CAT-ELEC", "Électricité & Câblage", "", "Matériel électrique et composants", "#2563EB"],
                ["CAT-ELEC-DISJ", "Disjoncteurs & Fusibles", "CAT-ELEC", "Disjoncteurs modulaires et différentiels", "#3B82F6"],
                ["CAT-QUIN", "Quincaillerie & Outillage", "", "Vis, chevilles et outils manuels", "#10B981"]
            ]
        },
        {
            "name": "CATALOGUE_PRODUITS",
            "headers": [
                "reference", "designation", "description", "code_barre",
                "code_categorie", "unite_mesure", "prix_achat_ht", "taux_tva",
                "prix_vente_ht", "prix_vente_ttc", "prix_vente_min",
                "stock_initial", "stock_minimum"
            ],
            "examples": [
                [
                    "DISJ-16A-SCH", "Disjoncteur 16A Courbe C Resi9", "Disjoncteur modulaire monophasé 16A Schneider", "3606480543210",
                    "CAT-ELEC-DISJ", "U", 38.50, 20.0,
                    65.00, 78.00, 55.00,
                    150, 20
                ],
                [
                    "CBL-3G15-100M", "Câble U1000 R2V 3G1.5mm² Couronne 100m", "Câble cuivre rigide d'alimentation 3G1.5", "3549210001452",
                    "CAT-ELEC", "ROULEAU", 240.00, 20.0,
                    360.00, 432.00, 310.00,
                    45, 10
                ]
            ]
        }
    ]

    for cfg in sheets_config:
        ws = wb.create_sheet(title=cfg["name"])
        headers = cfg["headers"]
        ws.append(headers)

        # Style des entêtes
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border
        
        # Lignes d'exemples
        for row_data in cfg["examples"]:
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if isinstance(val, (int, float)):
                    cell.alignment = align_right
                else:
                    cell.alignment = align_left

        # Figer la première ligne (entête)
        ws.freeze_panes = "A2"

        # Ajuster automatiquement la largeur des colonnes
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Supprimer la première feuille vide
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(output_filename)
    print(f"Fichier '{output_filename}' généré avec succès avec 4 feuilles : CLIENTS, FOURNISSEURS, CATEGORIES, CATALOGUE_PRODUITS.")

if __name__ == "__main__":
    build_onboarding_workbook()
